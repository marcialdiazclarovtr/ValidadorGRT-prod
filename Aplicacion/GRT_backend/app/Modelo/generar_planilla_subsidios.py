# generar_planilla_subsidios.py
# Llena la plantilla planilla_subsidios.xlsx con los datos de la validación

import shutil
from openpyxl import load_workbook

import logging
logger = logging.getLogger("uvicorn")

NOMBRE_HOJA = "Evidencia Subsidios 24 cuotas"

# Mapeo columna nombre → letra Excel (columnas desde B, A queda vacía)
COLUMNAS = {
    "Order Activity":     "B",
    "EQUIPMENT_RANK":     "C",
    "SKU":                "D",
    "OR":                 "E",
    "Cargo Activación":   "F",
    "Subsidio Claro up":  "G",
    "Subsidio 2 GRT":     "H",
    "Valor Full":         "I",
    "Valor full sin iva": "J",
    "Total Sin iva":      "K",
    "Total":              "L",
    "Resultado":          "M",
}

def generar_planilla_subsidios(
    template_path: str,
    output_path: str,
    resumen: list
):
    """
    Copia la plantilla a output_path y la llena con los datos del resumen.
    Cada fila del resumen corresponde a una combinación SKU + OR + Order Activity.
    Las columnas derivadas se escriben como fórmulas Excel.

    Columnas fórmulas (fila excel, base 4 porque filas 1-3 = encabezados):
      Valor full sin iva = Valor Full / 1.19                          → =I{fila}/1.19
      Total sin iva      = Valor full sin iva - Subs2 - SubsidioClaroUp → =J{fila}-H{fila}-G{fila}
      Total              = Total sin iva * 1.19 redondeado             → =ROUND(K{fila}*1.19,0)
      Resultado          = Total - Cargo Activación                    → =L{fila}-F{fila}
    """
    shutil.copy2(template_path, output_path)

    wb = load_workbook(output_path)

    if NOMBRE_HOJA not in wb.sheetnames:
        raise ValueError(f"No se encontró la hoja '{NOMBRE_HOJA}' en la plantilla.")

    ws = wb[NOMBRE_HOJA]

    # Descombinar todas las celdas antes de escribir
    # logger.info("Descombinando celdas")
    for merge in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merge))

    # logger.info(f"Celdas combinadas: {list(ws.merged_cells.ranges)}")

    fila_excel = 4  # Las filas 1-3 son encabezados en la plantilla

    for row in resumen:
        # Saltar filas de error de fecha (sku == "—")
        if row.get("sku") == "—":
            continue

        det = row.get("detalle_error") or {}

        # Valores directos
        ws[f"B{fila_excel}"] = row.get("order_activity") or ""
        ws[f"C{fila_excel}"] = ""  # EQUIPMENT_RANK siempre vacío
        ws[f"D{fila_excel}"] = row.get("sku") or ""
        ws[f"E{fila_excel}"] = row.get("or_val") or ""
        ws[f"F{fila_excel}"] = row.get("cargo_activacion") or ""
        ws[f"G{fila_excel}"] = row.get("subsidio_claroup") or ""
        ws[f"H{fila_excel}"] = row.get("subsidio_2_grt") or ""
        ws[f"I{fila_excel}"] = row.get("valor_full") or ""

        # Fórmulas Excel
        ws[f"J{fila_excel}"] = f"=I{fila_excel}/1.19"
        ws[f"K{fila_excel}"] = f"=J{fila_excel}-H{fila_excel}-G{fila_excel}"
        ws[f"L{fila_excel}"] = f"=ROUND(K{fila_excel}*1.19,0)"
        ws[f"M{fila_excel}"] = f"=L{fila_excel}-F{fila_excel}"

        fila_excel += 1  # Solo incrementa cuando se escribe una fila real

    wb.save(output_path)